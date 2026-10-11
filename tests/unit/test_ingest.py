from __future__ import annotations

import io
import json

import pytest

from enterprise_brain.ingest import (
    ALLOWED_SUFFIXES,
    IngestError,
    chunk_text,
    extract_text,
    safe_filename,
    sha256_bytes,
)

# ---------------------------------------------------------------- chunk_text


def test_chunk_empty_and_whitespace() -> None:
    assert chunk_text("") == []
    assert chunk_text("   \n\t  ") == []


def test_chunk_short_text_is_single_normalised_chunk() -> None:
    assert chunk_text("Hello   world.\n\nSecond  line.", size=100, overlap=10) == ["Hello world. Second line."]


def test_chunk_exactly_size_is_one_chunk() -> None:
    text = "a" * 100
    assert chunk_text(text, size=100, overlap=10) == [text]


def test_chunk_overlap_and_sentence_ends() -> None:
    text = "The quick brown fox jumps. " * 120
    chunks = chunk_text(text, size=600, overlap=120)
    assert len(chunks) > 1
    tail = chunks[0].split()[-4:]
    assert all(w in chunks[1].split() for w in tail)
    assert all(c.endswith((".", "!", "?")) for c in chunks)
    words = set(text.split())
    assert all(c.split()[0] in words for c in chunks)  # never begins mid-word


def test_chunk_boundary_at_limit_is_used() -> None:
    text = ("x" * 95) + " end. Next sentence follows here."
    assert chunk_text(text, size=100, overlap=10)[0] == ("x" * 95) + " end."


def test_chunk_sentence_lookahead_extends_end() -> None:
    text = ("x" * 95) + " endings now. Next sentence follows here and continues."
    assert chunk_text(text, size=100, overlap=10)[0] == ("x" * 95) + " endings now."


def test_chunk_lookahead_is_bounded() -> None:
    # the next sentence end is > 200 chars past the limit: snap to a word instead
    text = " ".join(["word"] * 120) + "."
    first = chunk_text(text, size=100, overlap=10)[0]
    assert len(first) <= 100
    assert not first.endswith(".")


def test_chunk_snaps_to_word_boundary_without_sentence() -> None:
    words = " ".join(f"word{i:03d}" for i in range(200))  # no punctuation at all
    chunks = chunk_text(words, size=100, overlap=20)
    assert len(chunks) > 2
    for c in chunks:
        for w in c.split():
            assert w.startswith("word") and len(w) == 7
        assert len(c) <= 100


def test_chunk_hard_split_when_no_spaces() -> None:
    chunks = chunk_text("y" * 250, size=100, overlap=20)
    assert chunks[0] == "y" * 100
    assert all(len(c) <= 100 for c in chunks)
    assert len(chunks) == 3


def test_chunk_snap_threshold_half_size() -> None:
    # the only space is at 30 (< size//2): no snap back, hard cut at 100
    text = "a" * 30 + " " + "b" * 200
    assert chunk_text(text, size=100, overlap=10)[0] == text[:100]


def test_chunk_snap_back_when_space_past_half() -> None:
    text = "a" * 60 + " " + "b" * 200
    assert chunk_text(text, size=100, overlap=10)[0] == "a" * 60


def test_chunk_snap_exactly_half_is_not_used() -> None:
    text = "a" * 50 + " " + "b" * 200
    assert chunk_text(text, size=100, overlap=10)[0] == text[:100]


def test_chunk_overlap_amount() -> None:
    text = " ".join(f"w{i:04d}" for i in range(300))  # 6 chars per token incl. space
    chunks = chunk_text(text, size=120, overlap=30)
    first, second = chunks[0].split(), chunks[1].split()
    shared = [w for w in first if w in second]
    assert 3 <= len(shared) <= 6
    assert second[0] == shared[0]


def test_chunk_covers_all_text() -> None:
    text = " ".join(f"token{i}" for i in range(500)) + "."
    chunks = chunk_text(text, size=200, overlap=40)
    assert set(" ".join(chunks).split()) == set(text.split())


@pytest.mark.parametrize(("size", "overlap"), [(0, 0), (100, 100), (100, 150), (100, -1)])
def test_chunk_rejects_bad_params(size: int, overlap: int) -> None:
    with pytest.raises(ValueError, match="size"):
        chunk_text("abc", size=size, overlap=overlap)


def test_chunk_zero_overlap_allowed() -> None:
    text = " ".join(f"t{i:03d}" for i in range(100))
    chunks = chunk_text(text, size=100, overlap=0)
    assert len(chunks) > 1
    assert not set(chunks[0].split()) & set(chunks[1].split())


# ---------------------------------------------------------------- extract_text


def test_extract_txt_and_md_with_bom() -> None:
    assert extract_text("\ufeffhello brain".encode(), ".txt") == "hello brain"
    assert extract_text(b"# Title\nbody", ".MD") == "# Title\nbody"


def test_extract_invalid_utf8_is_replaced() -> None:
    assert extract_text(b"ok \xff\xfe end", ".txt").startswith("ok ")


def test_extract_csv_rows_become_sentences() -> None:
    out = extract_text(b"client,revenue,region\nAcme,60000,\nGlobex,42500,EU\n\n", ".csv")
    assert out.splitlines() == ["client: Acme; revenue: 60000.", "client: Globex; revenue: 42500; region: EU."]


def test_extract_csv_header_only() -> None:
    assert extract_text(b"a,b\n", ".csv") == "a, b"


def test_extract_html_skips_scripts() -> None:
    html = (
        b"<html><head><style>x{}</style><script>alert(1)</script></head>"
        b"<body><h1>Title</h1><p>Body text</p></body></html>"
    )
    out = extract_text(html, ".html")
    assert "alert" not in out and "x{}" not in out
    assert out.index("Title") < out.index("Body text")
    assert "\n" in out


def test_extract_html_script_content_and_stray_close() -> None:
    # like browsers: script content is raw text up to the first </script>
    out = extract_text(b"</script>pre<script>a<script>b</script>c</script>visible", ".htm")
    assert out == "precvisible"
    assert extract_text(b"<style>p{}</style><noscript>n</noscript><template>t</template>ok", ".html") == "ok"


def test_extract_html_block_tags_break_lines() -> None:
    assert extract_text(b"<p>a</p><div>b</div><span>c</span>", ".html") == "\na\n\nb\nc"


def test_extract_json_flattens() -> None:
    data = {"client": "Acme", "contract": {"value": 240000, "status": "active"}, "tags": ["a", "b"]}
    out = extract_text(json.dumps(data).encode(), ".json").splitlines()
    assert out == ["client: Acme", "contract.value: 240000", "contract.status: active", "tags: a", "tags: b"]


def test_extract_json_scalar_and_list() -> None:
    assert extract_text(b'["x", 2]', ".json").splitlines() == ["x", "2"]


def test_extract_bad_json_raises() -> None:
    with pytest.raises(IngestError, match=r"\.json"):
        extract_text(b"{not json", ".json")


def test_extract_unsupported() -> None:
    with pytest.raises(IngestError, match="Unsupported file type: .exe"):
        extract_text(b"MZ", ".exe")
    with pytest.raises(IngestError, match=r"\(none\)"):
        extract_text(b"x", "")


def _pdf_bytes(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_extract_pdf() -> None:
    assert "Hello PDF world" in extract_text(_pdf_bytes("Hello PDF world"), ".pdf")


def test_extract_corrupt_pdf_raises() -> None:
    with pytest.raises(IngestError, match="Could not read .pdf file"):
        extract_text(b"%PDF-garbage", ".pdf")


def test_extract_docx() -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Contract value is 240000 dollars.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Vendor"
    table.rows[0].cells[1].text = "CloudServe"
    buf = io.BytesIO()
    d.save(buf)
    lines = extract_text(buf.getvalue(), ".docx").splitlines()
    assert "Contract value is 240000 dollars." in lines
    assert "Vendor | CloudServe" in lines


def test_allowed_suffixes() -> None:
    assert {".pdf", ".docx", ".txt", ".md", ".csv", ".html", ".htm", ".json"} == ALLOWED_SUFFIXES


# ---------------------------------------------------------------- helpers


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("report.pdf", "report.pdf"),
        ("../../etc/passwd", "passwd"),
        ("C:\\Users\\x\\file.txt", "file.txt"),
        ("we$ird<na>me?.md", "we_ird_na_me_.md"),
        ("   ", "upload"),
        ("...", "upload"),
        ("Q3 report (final).csv", "Q3 report _final_.csv"),
    ],
)
def test_safe_filename(raw: str, expected: str) -> None:
    assert safe_filename(raw) == expected


def test_safe_filename_truncates_keeping_extension() -> None:
    out = safe_filename("a" * 300 + ".pdf")
    assert len(out) == 200 and out.endswith(".pdf")
    assert safe_filename("b" * 300) == "b" * 200
    assert safe_filename("c" * 50 + ".txt", max_length=20) == "c" * 16 + ".txt"
    assert safe_filename("d" * 20, max_length=20) == "d" * 20


def test_sha256() -> None:
    assert sha256_bytes(b"abc") == "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"
