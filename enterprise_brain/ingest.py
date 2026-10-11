"""Ingest pipeline: extract text -> chunk (overlapping, sentence-snapped) -> embed -> store."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
from html.parser import HTMLParser
from pathlib import Path

from sqlalchemy import delete
from sqlalchemy.orm import Session

from enterprise_brain import models
from enterprise_brain.config import Settings
from enterprise_brain.llm import Embedder
from enterprise_brain.models import utcnow
from enterprise_brain.vectorstore import get_store

log = logging.getLogger(__name__)

ALLOWED_SUFFIXES = frozenset({".pdf", ".txt", ".md", ".csv", ".docx", ".html", ".htm", ".json"})
_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._ -]+")


class IngestError(ValueError):
    """A file could not be ingested (unsupported, empty, corrupt or too large)."""


def safe_filename(name: str, max_length: int = 200) -> str:
    """Strip any directory components and unsafe characters from a client filename."""
    base = name.replace("\\", "/").rsplit("/", 1)[-1].strip()
    base = _SAFE_NAME_RE.sub("_", base).strip(" .")
    if not base:
        base = "upload"
    if len(base) > max_length:
        stem, dot, ext = base.rpartition(".")
        base = (stem[: max_length - len(ext) - 1] + "." + ext) if dot else base[:max_length]
    return base


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


class _HTMLText(HTMLParser):
    _SKIP = frozenset({"script", "style", "noscript", "template"})
    _BLOCK = frozenset({"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section"})

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)


def _json_lines(value: object, prefix: str = "") -> list[str]:
    if isinstance(value, dict):
        out: list[str] = []
        for k, v in value.items():
            out.extend(_json_lines(v, f"{prefix}{k}." if prefix else f"{k}: "))
        return out
    if isinstance(value, list):
        out = []
        for item in value:
            out.extend(_json_lines(item, prefix))
        return out
    return [f"{prefix}{value}".strip()]


def extract_text(data: bytes, suffix: str) -> str:
    """Extract plain text from file bytes. Raises IngestError on bad input."""
    suffix = suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise IngestError(f"Unsupported file type: {suffix or '(none)'}")
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if suffix == ".docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            paras = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paras.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(paras)
        text = data.decode("utf-8-sig", errors="replace")
        if suffix in (".html", ".htm"):
            parser = _HTMLText()
            parser.feed(text)
            return "".join(parser.parts)
        if suffix == ".csv":
            rows = list(csv.reader(io.StringIO(text)))
            if len(rows) < 2:
                return "\n".join(", ".join(r) for r in rows)
            header = rows[0]
            # "col: value" rows make tabular facts retrievable as sentences.
            return "\n".join(
                "; ".join(f"{h}: {v}" for h, v in zip(header, row, strict=False) if v.strip()) + "."
                for row in rows[1:]
                if any(c.strip() for c in row)
            )
        if suffix == ".json":
            return "\n".join(_json_lines(json.loads(text)))
        return text
    except Exception as exc:  # corrupt PDFs/DOCX/JSON raise library-specific errors
        raise IngestError(f"Could not read {suffix} file: {exc}") from exc


def chunk_text(text: str, size: int = 600, overlap: int = 120) -> list[str]:
    """Split text into overlapping chunks of about ``size`` characters.

    A chunk end prefers a sentence boundary within a short lookahead past ``size``,
    otherwise it snaps back to a word boundary. Consecutive chunks overlap by about
    ``overlap`` characters so facts spanning a boundary stay retrievable, and every
    chunk after the first starts on a word boundary.
    """
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("require size > 0 and 0 <= overlap < size")
    text = " ".join(text.split())
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        if end < n:
            # a boundary may sit exactly at the limit, so look from end - 1
            m = re.search(r"[.!?](\s|$)", text[end - 1 : end + 200])
            if m:
                end += m.end() - 1
            else:
                snap = text.rfind(" ", start, end)
                if snap > start + size // 2:
                    end = snap
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= n:
            break
        start = max(end - overlap, start + 1)
        if text[start - 1] != " ":
            nxt = text.find(" ", start, end)
            if nxt != -1:
                start = nxt + 1
    return chunks


def store_document(
    db: Session,
    settings: Settings,
    embedder: Embedder,
    source: models.Source,
    filename: str,
    data: bytes,
    storage_path: str | None = None,
) -> models.Document:
    """Create or replace the Document ``filename`` in ``source`` and index its chunks.

    The caller commits. Raises IngestError when the file has no extractable text.
    """
    text = extract_text(data, Path(filename).suffix)
    pieces = chunk_text(text, settings.CHUNK_SIZE, settings.CHUNK_OVERLAP)
    if not pieces:
        raise IngestError(f"{filename}: no extractable text")
    vectors = embedder.embed(pieces)
    if len(vectors) != len(pieces):
        raise IngestError("embedding provider returned the wrong number of vectors")

    doc = (
        db.query(models.Document)
        .filter(models.Document.source_id == source.id, models.Document.filename == filename)
        .one_or_none()
    )
    if doc is None:
        doc = models.Document(source_id=source.id, filename=filename)
        db.add(doc)
    else:
        db.execute(delete(models.Chunk).where(models.Chunk.document_id == doc.id))
    doc.sha256 = sha256_bytes(data)
    doc.size_bytes = len(data)
    doc.storage_path = storage_path
    doc.chunk_count = len(pieces)
    doc.updated_at = utcnow()
    db.flush()

    store = get_store(settings)
    for position, (piece, vector) in enumerate(zip(pieces, vectors, strict=True)):
        chunk = models.Chunk(
            source_id=source.id,
            document_id=doc.id,
            position=position,
            text=piece,
            embedding_json=json.dumps(vector),
            embedding_model=embedder.model_id,
            embedding_dim=len(vector),
        )
        db.add(chunk)
        db.flush()
        store.upsert(db, chunk.id, vector)
    log.info("document indexed", extra={"source_id": source.id, "document": filename, "chunks": len(pieces)})
    return doc


def reindex_all(db: Session, settings: Settings, embedder: Embedder, batch: int = 128) -> int:
    """Re-embed every chunk with ``embedder`` (after changing EMBEDDINGS). Returns count."""
    store = get_store(settings)
    total = 0
    last_id = 0
    while True:
        rows = (
            db.query(models.Chunk)
            .filter(models.Chunk.id > last_id)
            .order_by(models.Chunk.id)
            .limit(batch)
            .all()
        )
        if not rows:
            break
        vectors = embedder.embed([c.text for c in rows])
        for chunk, vector in zip(rows, vectors, strict=True):
            chunk.embedding_json = json.dumps(vector)
            chunk.embedding_model = embedder.model_id
            chunk.embedding_dim = len(vector)
            store.upsert(db, chunk.id, vector)
        db.commit()
        total += len(rows)
        last_id = rows[-1].id
    return total
